#!/usr/bin/env python3
"""Generate the account strategy presentation and its evidence artefacts.

Run from any directory: python scripts/generate_presentation.py
"""
from __future__ import annotations

import csv
import json
import re
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "presentations"
AS_OF_ISO = "2026-09-16"
PPTX = OUT / f"M365_Copilot_License_Expansion_ISD_Strategy_{AS_OF_ISO.replace('-', '')}.pptx"
MATRIX = OUT / "account_evidence_matrix.csv"
MANIFEST = OUT / "source_manifest.json"
AS_OF = AS_OF_ISO.replace("-", "年", 1).replace("-", "月", 1) + "日"
EXPECTED_ACCOUNT_COUNT = 52
EXPECTED_INTRO_SLIDES = 23
W, H = Inches(13.333), Inches(7.5)
NAVY, BLUE, TEAL, LIGHT, MID, RED, AMBER, GREEN, WHITE, GRAY = (
    "17365D", "0078D4", "00A6A6", "F4F7FA", "D9E2F3", "C50F1F",
    "D83B01", "107C10", "FFFFFF", "5B6573",
)

# These are analytical classifications, not source claims. All unlisted accounts remain
# intentionally qualified rather than promoted on the basis of allocation size.
PROFILES = {
    "トヨタ自動車 TMC": ("営業優先", "利用待ち・研修・需要確認", "購買済/追加/無償の基準が混在",
                        "需要台帳→部門導入→定着計測→追加数量の承認", "高"),
    "トヨタテクニカルディベロップメント株式会社": ("営業優先", "全社展開/希望者購入の運用化", "全社利用可能=全員購入ではない",
                        "対象者公募→説明会→継続利用→部門別購入判断", "高"),
    "ジェイテクト": ("営業優先", "部門別公募と評価設計", "購入時期・決裁・成功基準が未確認",
                        "候補部門→ユースケース評価→継続根拠→次年度購入", "中"),
    "トヨタファイナンス": ("営業優先", "配布後研修と金融業務の評価", "予算/有償継続は未確認",
                        "対象コホート→研修→反復利用→有償継続決裁", "中"),
    "トヨタ自動車東日本": ("営業優先", "必要数量の試行・効果測定", "TTS/MS記述と展開時点を要照合",
                        "ガードレール→試行→効果確認→必要数量承認", "中"),
    "スバル Subaru": ("条件付き", "権限・データ・配布設計", "価格/急な全社E7化に懸念",
                       "限定グループの安全設計→利用証跡→限定拡大判断", "高"),
    "AISIN 本体": ("条件付き", "グループ標準化", "一律E7化・価格・公平性",
                 "申請/配布基準→共通研修→利用可視化→有償標準化判断", "高"),
    "トヨタ車体 Toyota Shatai": ("条件付き", "段階導入と個別業務の具体化", "数量確約とユースケースが未成熟",
                         "500想定の対象選定→業務実証→段階購入判定", "中"),
    "豊田通商 TTC": ("条件付き", "本体の成功事例化", "ITコスト・会社別予算・既存研修",
                 "既存施策補完→活用証跡→関連会社別の採否判断", "中"),
    "トヨタ紡織 Toyota Boshoku 本体": ("条件付き", "Agent/Copilot価値・ガバナンス", "価格・必要性・数量コミット",
                              "少数ユースケース→統制/効果→価格を含む継続判断", "高"),
    "マツダ Mazda": ("別製品優先", "Copilot Studio/Agentの価値検証", "Chatで十分、R&D配布停止/開始記述が競合",
                  "Agent候補→利用量/成果→Studio/Agent365の採否", "高"),
    "トヨタコニックプロ": ("条件付き", "更新時のCopilot/Agent評価", "価格理解が前提",
                     "更新要件→小規模評価→更新/追加の判断", "中"),
    "TMHE": ("条件付き", "約300人展開の組織変革", "9,000無償は注文ではない",
             "PoC対象→変革支援→継続利用→親会社方針と購入確認", "中"),
    "CFAO": ("条件付き", "既存外部支援と非重複の定着支援", "予算限定・全社計画なし",
             "既存支援範囲確認→限定導入→成果確認→購入可否", "中"),
    "Bastian": ("要確認", "成長意向の適格性確認", "購入規模/課題/決裁が不足",
                "決裁者・対象・予算確認→必要時のみ導入支援", "低"),
    "VANDERLANDE": ("別製品優先", "E5移行要件の評価", "価格/評価懸念、Copilot感度は限定的",
                  "E5要件確認→既存権利確認→必要時のみ追加提案", "中"),
    "TMCA": ("条件付き", "限定E7の商用条件確認", "価格条件、Agent365必要性なし",
             "対象限定→商用条件→必要性確認→購入判断", "中"),
    "DENSO 本体": ("保留", "改善後のROI・対象者整理", "割引/契約条件の改善が前提、競合も評価",
               "条件改善→限定シナリオ→価値証拠→社内決裁可否", "高"),
    "愛三工業": ("保留", "現行環境の安定化", "旧来の拡大期待より現環境課題が優先",
             "環境課題整理→利用可能性再確認→必要時に評価", "中"),
    "TOYOTA MOTOR EUROPE": ("保留", "既存投資・利用低下の確認", "投資保留/他AI・地域要件",
                         "利用/投資方針確認→再評価", "中"),
}


def rgb(value):
    return RGBColor.from_string(value)


def clean(text):
    return re.sub(r"\s+", " ", text).strip()


def lines_for(path):
    return path.read_text(encoding="utf-8").splitlines()


def field(lines, prefix, default):
    return next((x.split(":", 1)[1].strip() for x in lines if x.startswith(prefix)), default)


def accounts():
    result = []
    for path in sorted(DOCS.glob("*.md")):
        if path.name == "00_INDEX.md":
            continue
        lines = lines_for(path)
        title = next((x[2:].strip() for x in lines if x.startswith("# ")), path.stem)
        region = field(lines, "- リージョン:", "不明")
        # The source has two distinct documents headed simply "TFS"; keep them
        # unambiguous in the evidence matrix and appendix.
        if title == "TFS":
            title = f"TFS {region}"
        usage = field(lines, "- 総ユーザー数:", "記載なし")
        licenses = field(lines, "- Copilot 有償:", "記載なし")
        # Keep a source excerpt verbatim (apart from whitespace) and preserve truncation marks.
        start = next((i for i, x in enumerate(lines) if x.startswith("### 2026/09")), None)
        if start is None or any("記載なし" in x for x in lines[start + 1:start + 4]):
            start = next((i for i, x in enumerate(lines) if x.startswith("### 2026/07")), None)
        source_date = re.sub(r"^###\s*", "", lines[start]).strip() if start is not None else "日付不明"
        excerpt, refs = [], []
        for i in range((start + 1) if start is not None else len(lines), len(lines)):
            x = clean(lines[i])
            if x.startswith("### ") or x.startswith("## ") or x.startswith("---"):
                if excerpt:
                    break
                continue
            if x and not x.startswith("_記載なし_") and not x.startswith("出典:"):
                excerpt.append(x)
                refs.append(i + 1)
            if len(" ".join(excerpt)) >= 120:
                break
        evidence = " ".join(excerpt)[:260] or "顧客フィードバックの具体記載なし"
        line_ref = f"{min(refs)}–{max(refs)}" if refs else "該当なし"
        profile = PROFILES.get(title, ("要確認", "購入意思・障害の確認", "資料の根拠が限定的",
                                       "決裁者・予算・対象・成功基準を確認→必要時に限定支援", "低"))
        result.append({
            "name": title, "path": f"docs/{path.name}", "region": region, "usage": usage,
            "licenses": licenses, "evidence": evidence, "lines": line_ref,
            "source_date": source_date,
            "decision": profile[0], "isd": profile[1], "barrier": profile[2],
            "chain": profile[3], "confidence": profile[4],
        })
    return result


def add_box(slide, x, y, w, h, text="", fill=WHITE, line=WHITE, size=14, color=NAVY,
            bold=False, valign=MSO_ANCHOR.TOP, margin=0.12):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    tf = shape.text_frame; tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.text = text; p.font.name = "Noto Sans CJK JP"
    p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color)
    p.alignment = PP_ALIGN.LEFT
    return shape


def footer(slide, n, source="分析・出典は各スライド記載"):
    add_box(slide, Inches(.35), Inches(7.13), Inches(12.6), Inches(.20),
            f"{source}  |  {AS_OF}  |  顧客内部情報を含む／取扱注意", fill=WHITE, line=WHITE, size=7, color=GRAY, margin=0)
    add_box(slide, Inches(12.75), Inches(7.07), Inches(.22), Inches(.22), str(n),
            fill=NAVY, line=NAVY, size=7, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE, margin=0)


def title(slide, text, subtitle="", n=1, source=""):
    add_box(slide, Inches(.35), Inches(.28), Inches(12.6), Inches(.62), text, fill=NAVY, line=NAVY,
            size=25, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_box(slide, Inches(.48), Inches(.96), Inches(12), Inches(.34), subtitle, fill=WHITE, line=WHITE,
                size=10, color=GRAY)
    footer(slide, n, source or "分析資料：docs 配下の52アカウント記録")


def bullets(slide, items, x=.55, y=1.45, w=12.1, h=5.35, size=16):
    shape = add_box(slide, Inches(x), Inches(y), Inches(w), Inches(h), fill=LIGHT, line=LIGHT, size=size)
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = "• " + item; p.font.name = "Noto Sans CJK JP"; p.font.size = Pt(size)
        p.font.color.rgb = rgb(NAVY); p.level = 0; p.space_after = Pt(10)


def main_slide(prs, heading, subtitle, items, source="分析・推薦（資料記載と区別）"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    n = len(prs.slides); title(slide, heading, subtitle, n, source); bullets(slide, items)


def appendix_slide(prs, a):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); n = len(prs.slides)
    title(slide, f"アカウント詳細｜{a['name']}", f"{a['region']}  |  資料内日付: {a['source_date']}  |  判断: {a['decision']}  |  根拠信頼度: {a['confidence']}", n,
          f"{a['path']}:{a['lines']}（GitHub commit は source_manifest.json を参照）")
    add_box(slide, Inches(.45), Inches(1.42), Inches(3.9), Inches(1.35),
            f"台帳値（活動実績ではない）\n総ユーザー: {a['usage']}\nCopilot 有償/無償/合計: {a['licenses']}", fill=LIGHT, line=MID, size=11, bold=True)
    add_box(slide, Inches(4.55), Inches(1.42), Inches(8.3), Inches(1.35),
            f"資料の抜粋（事実/見解は原文）\n{a['evidence']}", fill=WHITE, line=MID, size=11)
    add_box(slide, Inches(.45), Inches(3.02), Inches(5.95), Inches(1.35),
            f"商用ゲート・障壁\n{a['barrier']}\n次の確認: 購入決裁者、予算時期、対象コホート、既存権利、成功基準", fill="FFF4CE", line=AMBER, size=11)
    add_box(slide, Inches(6.65), Inches(3.02), Inches(6.2), Inches(1.35),
            f"ISD適合（提案/推論）\n{a['isd']}\n根拠: {'資料に支援ニーズの記載あり' if a['decision'] not in ('要確認', '保留') else '資料上は仮説。確認後に開始'}", fill="E8F5F5", line=TEAL, size=11)
    add_box(slide, Inches(.45), Inches(4.65), Inches(12.4), Inches(1.35),
            f"支援 → 運用結果 → 購入/利用拡大（推薦）\n{a['chain']}\n販売優先度は意思・障壁・根拠に基づく定性判断であり、売上額・確率・活動率を示すものではない。", fill="EAF3F8", line=BLUE, size=12)


def write_artefacts(data, commit, appendix_start):
    OUT.mkdir(exist_ok=True)
    with MATRIX.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["account", "region", "source_date", "source_path", "source_lines", "source_excerpt",
                                                "license_header", "decision_label", "commercial_barrier", "isd_proposal",
                                                "support_to_purchase_chain", "confidence", "appendix_slide"])
        writer.writeheader()
        for i, a in enumerate(data, start=appendix_start):
            writer.writerow({"account": a["name"], "region": a["region"], "source_date": a["source_date"], "source_path": a["path"], "source_lines": a["lines"],
                             "source_excerpt": a["evidence"], "license_header": a["licenses"], "decision_label": a["decision"],
                             "commercial_barrier": a["barrier"], "isd_proposal": a["isd"], "support_to_purchase_chain": a["chain"],
                             "confidence": a["confidence"], "appendix_slide": i})
    MANIFEST.write_text(json.dumps({"as_of": AS_OF_ISO, "analyzed_commit": commit, "index": "docs/00_INDEX.md",
                                    "account_count": len(data), "documents": [{"path": a["path"]} for a in data],
                                    "limitations": ["台帳の有償/無償/合計は活動ユーザー・新規受注ではない。",
                                                    "原資料には不整合、欠損、または切れた記述（[...]）があり、補完していない。",
                                                    "価格、包装、権利、契約条件は営業が確認する。"]},
                                   ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def validate(data):
    with zipfile.ZipFile(PPTX) as z:
        bad = z.testzip()
        if bad is not None:
            raise ValueError(f"corrupt ZIP member: {bad}")
        slides = [x for x in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", x)]
        if len(slides) != len(data) + EXPECTED_INTRO_SLIDES:
            raise ValueError(f"unexpected slide count: {len(slides)}")
        for name in slides:
            ET.fromstring(z.read(name))
    reopened = Presentation(PPTX)
    if len(reopened.slides) != len(data) + EXPECTED_INTRO_SLIDES or len(data) != EXPECTED_ACCOUNT_COUNT:
        raise ValueError("slide or account count validation failed")
    if not all(a["name"] for a in data):
        raise ValueError("empty account name")
    if any("TODO" in shape.text for slide in reopened.slides for shape in slide.shapes if hasattr(shape, "text")):
        raise ValueError("unresolved placeholder found")


def generate():
    data = accounts()
    if len(data) != EXPECTED_ACCOUNT_COUNT:
        raise ValueError(f"index scope mismatch: {len(data)} account files")
    try:
        commit = subprocess.check_output(["git", "-C", str(ROOT), "rev-parse", "HEAD"], text=True, stderr=subprocess.DEVNULL).strip()
    except subprocess.CalledProcessError:
        commit = "unknown (not generated from a git checkout)"
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    main_slide(prs, "Microsoft ライセンス拡大とISD支援機会", f"{AS_OF}｜52アカウントの資料統合｜顧客内部情報を含む／取扱注意",
               ["目的：実際のMicrosoft 365 Copilot利用定着を起点に、有償ライセンスの継続・拡大機会を見極める。",
                "資料根拠と推薦を区別し、規模・無償枠・課題の大きさだけで優先順位を上げない。", "対象資料：docs/00_INDEX.md および全52アカウント文書。"])
    main_slide(prs, "エグゼクティブ提言", "営業優先、ISD適合、条件付き商材を分離",
               ["営業優先：TMC、TTDC、ジェイテクトJP、トヨタファイナンス、トヨタ自動車東日本。まず購入意向・数量・決裁を検証する。",
                "ISDの具体的な入口：TTDC（展開/研修）、Subaru（安全な配布設計）、金融/東日本（試行・効果測定）、TMC（需要/定着）。",
                "条件付き：Subaru、AISIN、Toyota Shatai、TTC。E7/Agent等は価格、対象限定、既存権利、必要性を確認後に判断。"])
    main_slide(prs, "評価方法と商用ゲート", "資料記載（事実/顧客見解/TTS・MS見解）と推薦を分離",
               ["販売優先度＝明示された購入/拡大意向、時期、商用ゲート、根拠の新しさで定性評価。ISD課題の大きさや配布枠だけでは上げない。",
                "ISD開始前に、購入決裁者・予算時期・対象コホート・既存支援/権利・成功基準を確認する。",
                "信頼度：高=直近かつ具体、 中=具体だが未確定/競合、低=関心のみ又は詳細不足。"])
    main_slide(prs, "利用・ライセンス・収益を混同しない", "台帳ヘッダーは活動実績・追加受注ではない",
               ["有償/無償/合計は資料ヘッダー値。割当、活動、継続、有償転換、新規増分を別の指標として扱う。",
                "グループ/本社の重複、総ユーザー0/欠損、Vanderlandeの合計超過等があるため、横断浸透率・売上合計は算出しない。",
                "E7切替は置換分を差し引く。実価格、包装、権利をこの資料の外部検証済み事実として記載しない。"])
    main_slide(prs, "ポートフォリオ：定性マップ", "横軸=商用準備度、縦軸=ISDで解ける導入障害。金額・確率ではない",
               ["右上：TTDC、Subaru、トヨタファイナンス、東日本、TMC（ただし需要/数量を再照合）。",
                "右下/条件付き：AISIN、Toyota Shatai、TTC、Toyota Boshoku、TMHE、CFAO。",
                "左側（商用ゲート先行）：DENSO、Mazda、TME、Vanderlande、TMCA。支援ニーズは購買意思の証明ではない。"])
    main_slide(prs, "優先アカウント：Copilot拡大", "購入意向と導入阻害が両方確認できる順に検証",
               ["TMC：待機・追加購入の記述はあるが、TTSとMSで人数/基準が異なる。重複計上せず、需要台帳を確定する。",
                "TTDC：全社展開/希望者購入の記載。利用可能化を全員購入と解釈せず、対象部門別の継続利用を確認。",
                "ジェイテクトJP：部門別希望と次年度購入への前向きさを、評価設計と購買ゲートに結び付ける。"])
    main_slide(prs, "優先アカウント：導入定着を購入判断へ", "ISDは一般セミナーではなく、明確な障害除去に限定",
               ["トヨタファイナンス：配布と研修調整はあるが、予算/有償継続は未確認。研修後の反復利用と決裁を出口にする。",
                "トヨタ自動車東日本：必要量、試行、効果測定の記述をTTS/MS間で照合。DLP/ガードレールを先に整える。",
                "Subaru：E3+Copilotの拡大関心。限定グループへのE7を含む選択肢を、価格とアクセス権課題の解消後に検討。"])
    for name, focus in [("TTDC", "対象者公募、説明会、部門別ユースケース、反復利用の可視化、継続購入の決裁"),
                        ("Subaru", "データ/権限/割当の安全設計、限定導入、利用証跡、対象拡大判断"),
                        ("トヨタファイナンス", "金融業務コホート、研修、日常業務での利用確認、有償継続の決裁"),
                        ("トヨタ自動車東日本", "DLP/ガバナンス、短期試行、効果測定、必要数量の承認"),
                        ("ジェイテクトJP", "部門公募、評価指標、次年度予算に向けた根拠づくり"),
                        ("TMC", "待機者・購入済・無償を照合、部門別オンボーディング、追加需要の確認")]:
        main_slide(prs, f"ISD介入チェーン｜{name}", "提案：支援の出口を「購入または継続の判断」に定義",
                   [f"入口：{focus}", "役割分担：営業=商用条件/決裁、顧客IT=対象・統制、ISD=設計/定着/測定。", "出口：活動/反復利用の定義、試行完了、根拠に基づく継続・追加・見送り判断。時間短縮を自動的な現金効果とみなさない。"])
    main_slide(prs, "条件付き：E7・Agent365・E5・Studio", "製品名は資料内の顧客仮説/用語。権利・価格・包装は営業確認",
               ["Toyota Boshoku：約10,000 E7は条件付き言及で、5,000–6,000はエージェント数。人/活動ユーザー/購入確約ではない。",
                "Mazda：Copilot Chatで十分との見解があり、Agent開発/消費の条件付きルート。R&D 5,000の配布停止/開始記述は確認する。",
                "Vanderlande：E5移行関心が先行。既存E5権利で満たせる統制要件は追加販売と直結させない。"])
    main_slide(prs, "国際機会と既存パートナー", "規模ではなく商用準備度と支援の差分で選別",
               ["TMHE：好意的PoC/約300展開の記載。9,000無償配布は注文ではなく、組織変革を限定支援する。",
                "CFAO：成長関心と外部コンサル利用。予算・全社計画を確認し、既存パートナーと重ならない定着支援のみ提案。",
                "Bastian：拡大意向の詳細が限定的。ISD投資前に決裁者、課題、予算、対象を適格化する。"])
    main_slide(prs, "優先度を下げる/保留するケース", "長期機会を否定せず、直近の拡大案件とは区別",
               ["DENSO：改善された割引/契約条件が前提で、Copilot採用ありきではない。ROI支援要望は予算・購入意向の証明ではない。",
                "Mazda：広範なM365 Copilot座席の前提を置かず、Agent条件を確認。", "愛三：9月は現行環境優先。Kayaba：ROI停滞、TME：投資保留/利用低下。購入意思を再確認する。"])
    main_slide(prs, "提案する30/60/90日アクション", "相対的な提案であり、確定した顧客プロジェクト日程ではない",
               ["30日：各優先アカウントで決裁者・予算・対象・既存権利・成功基準を記録し、TMC/Toyota Eastの数値を照合。",
                "60日：安全な対象コホートでオンボーディング/ユースケース/測定を実施（営業、顧客IT、ISDの役割分担）。",
                "90日：継続利用と業務結果をレビューし、追加購入・更新・限定Upsell・見送りを根拠付きで判定。"])
    main_slide(prs, "KPIとGo/No-Go", "数値目標は資料にないため、定義を合意して追跡する",
               ["割当ユーザー、活動ユーザー、反復利用ユーザー（定義を合意）、試行完了、業務結果の証拠、有償継続承認。",
                "商用：増分有償席数、置換契約額と純増、決裁日。架空の金額/勝率/ROIは置かない。",
                "Go：対象/安全性/成功基準/スポンサーが揃う。No-Go/保留：価格・契約・必要性・予算のゲート未解決。"])
    main_slide(prs, "オープンクエスチョンとデータ照合", "次回アカウントオーナー会話で確定する項目",
               ["TMC：TTS約1,500とMS待機1万、購入済/追加1万と既存カウントの関係。", "Mazda：R&D 5,000の停止/承認・開始状況。Toyota Boshoku：Agent数と人/活動/契約数の区別。",
                "全件共通：台帳値の時点・割当・活動・有償転換、グループ重複、購入決裁/予算/既存権利。期限前の予定は実施済みと扱わない。"])
    main_slide(prs, "営業優先の運用ルール", "案件化前に「売れる見込み」と「支援が必要」を分離",
               ["営業：顧客の購入意思、対象人数、予算・決裁日、商用条件を文書化する。", "ISD：障害が導入設計・安全性・活用定着・測定である場合に、限定スコープで支援する。",
                "共通：価格交渉のみ、必要性なし、既存支援との重複、活動根拠なしの場合は支援を案件化しない。"])
    main_slide(prs, "データ品質と競合する記述の扱い", "結論を強めず、資料内の最新日付・原文・不確実性を表示",
               ["9月フィードバックを優先する。ただし7月見出し内に9月の日付がある場合は実際の日付を確認する。",
                "予定された会議/トライアルは、完了を明示する記述がない限り実施済みとはしない。", "欠損「-」、総数0、矛盾した合計、切れた記述はそのまま制限事項として残す。"])
    main_slide(prs, "商材別の適用ガードレール", "顧客の用語・仮説を、確定した製品提案へ読み替えない",
               ["M365 Copilot：日常業務の反復利用・実務成果・有償継続の根拠を優先。", "Agent365/Copilot Studio：業務エージェントの必要性、統制、利用量/成果、商用条件を確認。",
                "E5/E7：セキュリティ/統制要件と既存権利を確認し、切替の置換分を考慮して純増を評価。"])
    main_slide(prs, "付録の読み方と出典管理", "全52アカウントを1枚ずつ掲載。詳細な追跡はCSV/manifestで行う",
               ["各カードは資料内日付、台帳ヘッダー、短い原文抜粋、障壁、推薦、次の確認を表示する。", "原文のファイル・行範囲は各スライド下部に表示。全行・コミット・全件対応は account_evidence_matrix.csv と source_manifest.json に記録。",
                "本資料はMicrosoft公式の推奨・価格表・商用コミットメントではない。"])
    if len(prs.slides) != EXPECTED_INTRO_SLIDES:
        raise ValueError(f"unexpected intro slide count: {len(prs.slides)}")
    write_artefacts(data, commit, len(prs.slides) + 1)
    for a in data:
        appendix_slide(prs, a)
    prs.save(PPTX)
    validate(data)
    print(f"Generated {PPTX} ({len(prs.slides)} slides), matrix ({len(data)} accounts), manifest at commit {commit}")


if __name__ == "__main__":
    generate()
